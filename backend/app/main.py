from io import BytesIO
import re
from fastapi import FastAPI, UploadFile, File, Depends, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from .database import Base, engine, get_db
from .models import Lead, Campaign, TelegramAccount
from .schemas import LeadOut, CampaignCreate

Base.metadata.create_all(bind=engine)

app = FastAPI(title="Infinity Sales API", version="0.1.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def normalize_phone(value: str) -> str:
    digits = re.sub(r"\D", "", value or "")
    if not digits:
        return ""
    if digits.startswith("998"):
        return "+" + digits
    if len(digits) == 9:
        return "+998" + digits
    return "+" + digits

@app.get("/health")
def health():
    return {"status": "ok", "version": "0.1.0"}

@app.get("/dashboard")
def dashboard(db: Session = Depends(get_db)):
    total = db.query(Lead).count()
    telegram_found = db.query(Lead).filter(Lead.telegram_status == "found").count()
    contacted = db.query(Lead).filter(Lead.status.in_(["contacted","replied","interested","negotiation","won"])).count()
    replied = db.query(Lead).filter(Lead.status.in_(["replied","interested","negotiation","won"])).count()
    return {
        "total_leads": total,
        "telegram_found": telegram_found,
        "contacted": contacted,
        "replied": replied,
    }

@app.get("/leads", response_model=list[LeadOut])
def list_leads(limit: int = 200, db: Session = Depends(get_db)):
    return db.query(Lead).order_by(Lead.id.desc()).limit(min(limit, 1000)).all()

@app.post("/leads/import")
async def import_leads(file: UploadFile = File(...), db: Session = Depends(get_db)):
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "For MVP use an .xlsx file")
    raw = await file.read()
    wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"imported": 0, "skipped": 0}

    headers = [str(x or "").strip().lower() for x in rows[0]]
    aliases = {
        "company": ["company","компания","название","организация"],
        "first_name": ["first_name","имя","name"],
        "last_name": ["last_name","фамилия","surname"],
        "phone": ["phone","телефон","номер"],
        "city": ["city","город"],
        "category": ["category","категория","тип"],
        "instagram": ["instagram","инстаграм"],
        "website": ["website","site","сайт"],
        "comment": ["comment","комментарий","примечание"],
    }
    idx = {}
    for key, variants in aliases.items():
        for v in variants:
            if v in headers:
                idx[key] = headers.index(v)
                break

    if "phone" not in idx:
        raise HTTPException(400, "Не найдена колонка phone/телефон/номер")

    imported = 0
    skipped = 0
    for row in rows[1:]:
        phone = normalize_phone(str(row[idx["phone"]] or ""))
        if not phone:
            skipped += 1
            continue
        if db.query(Lead).filter(Lead.phone == phone).first():
            skipped += 1
            continue

        def val(name):
            i = idx.get(name)
            return str(row[i] or "").strip() if i is not None and i < len(row) else ""

        db.add(Lead(
            company=val("company"),
            first_name=val("first_name"),
            last_name=val("last_name"),
            phone=phone,
            city=val("city"),
            category=val("category"),
            instagram=val("instagram"),
            website=val("website"),
            comment=val("comment"),
        ))
        imported += 1
    db.commit()
    return {"imported": imported, "skipped": skipped}

@app.post("/campaigns")
def create_campaign(payload: CampaignCreate, db: Session = Depends(get_db)):
    item = Campaign(name=payload.name, product=payload.product, template=payload.template)
    db.add(item)
    db.commit()
    db.refresh(item)
    return {"id": item.id, "status": item.status}

@app.get("/campaigns")
def list_campaigns(db: Session = Depends(get_db)):
    return db.query(Campaign).order_by(Campaign.id.desc()).all()

@app.get("/telegram/accounts")
def telegram_accounts(db: Session = Depends(get_db)):
    return db.query(TelegramAccount).all()

@app.get("/inbox")
def inbox():
    return {"items": [], "note": "Telegram inbox sync lands in MVP 0.2"}
